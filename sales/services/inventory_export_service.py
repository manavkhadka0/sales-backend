import io
from typing import Any, Dict, List

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class InventoryChangeLogExportService:
    """
    Service responsible for exporting inventory change logs (Add and Update actions only)
    into formatted Excel (.xlsx) downloads.
    """

    EXCEL_HEADERS: List[str] = [
        "Date",
        "Product Name",
        "Action",
        "Old Quantity",
        "New Quantity",
        "Difference",
        "Franchise",
        "Changed By",
    ]

    @classmethod
    def get_log_row_data(cls, log) -> Dict[str, Any]:
        diff = log.new_quantity - log.old_quantity
        diff_str = f"+{diff}" if diff > 0 else str(diff)

        inv = log.inventory
        prod_name = inv.product.name if inv and inv.product else "N/A"

        franchise_name = "N/A"
        if inv and inv.franchise:
            franchise_name = inv.franchise.name
        elif inv and inv.distributor:
            franchise_name = f"Distributor: {inv.distributor.name}"
        elif inv and inv.factory:
            franchise_name = f"Factory: {inv.factory.name}"

        user_name = "System"
        if log.user:
            user_name = (
                log.user.get_full_name()
                or log.user.first_name
                or log.user.email
                or "Unknown"
            )

        changed_at_str = (
            log.changed_at.strftime("%Y-%m-%d %H:%M:%S") if log.changed_at else "N/A"
        )
        action_str = (
            log.get_action_display()
            if hasattr(log, "get_action_display")
            else str(log.action).capitalize()
        )

        return {
            "date": changed_at_str,
            "product_name": prod_name,
            "action": action_str,
            "old_quantity": log.old_quantity,
            "new_quantity": log.new_quantity,
            "difference": diff_str,
            "franchise": franchise_name,
            "changed_by": user_name,
        }

    @classmethod
    def export_to_excel(cls, queryset) -> HttpResponse:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Logs"

        # Ensure grid lines are visible
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        # Write Headers
        ws.append(cls.EXCEL_HEADERS)
        for col_idx in range(1, len(cls.EXCEL_HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        # Zebra striping fills
        even_fill = PatternFill(
            start_color="F2F4F7", end_color="F2F4F7", fill_type="solid"
        )

        # Write Data Rows
        row_num = 2
        for log in queryset.iterator(chunk_size=1000):
            data = cls.get_log_row_data(log)
            row_values = [
                data["date"],
                data["product_name"],
                data["action"],
                data["old_quantity"],
                data["new_quantity"],
                data["difference"],
                data["franchise"],
                data["changed_by"],
            ]
            ws.append(row_values)

            # Apply cell formatting
            for col_idx in range(1, len(row_values) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border

                if row_num % 2 == 0:
                    cell.fill = even_fill

                # Align numbers right, actions and date center, names left
                if col_idx in [4, 5, 6]:
                    cell.alignment = right_align
                elif col_idx in [1, 3]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

            ws.row_dimensions[row_num].height = 20
            row_num += 1

        # Adjust Column Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Save to buffer
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="inventory_changelog_export_{timestamp}.xlsx"'
        )
        return response
